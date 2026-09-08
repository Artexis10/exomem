"""Generate synthetic document/OCR fixtures without checked-in binary blobs."""

from __future__ import annotations

import shutil
import sys
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import fitz
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation


def _font() -> ImageFont.FreeTypeFont:
    return ImageFont.load_default(size=42)


def _write_image(path: Path, text: str) -> None:
    image = Image.new("RGB", (1050, 180), "white")
    ImageDraw.Draw(image).text((30, 55), text, fill="black", font=_font())
    image.save(path)


def _write_docx(path: Path) -> None:
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>""",
        )
        archive.writestr(
            "_rels/.rels",
            """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>""",
        )
        archive.writestr(
            "word/document.xml",
            """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>
  <w:p><w:r><w:t>DOCX TABLE HEADING</w:t></w:r></w:p>
  <w:tbl>
    <w:tr><w:tc><w:p><w:r><w:t>Item</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>Quantity</w:t></w:r></w:p></w:tc></w:tr>
    <w:tr><w:tc><w:p><w:r><w:t>Widget</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>3</w:t></w:r></w:p></w:tc></w:tr>
  </w:tbl>
  <w:sectPr><w:pgSz w:w="12240" w:h="15840"/></w:sectPr>
</w:body></w:document>""",
        )


def generate(root: Path) -> None:
    root.mkdir(parents=True, exist_ok=True)

    image_path = root / "image-ocr.png"
    _write_image(image_path, "IMAGE OCR 42")

    scan_path = root / "scanned-page.png"
    _write_image(scan_path, "SCANNED PDF BRAVO 73")
    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "PDF SEARCHABLE ALPHA", fontsize=20)
    page = pdf.new_page(width=1050, height=180)
    page.insert_image(page.rect, filename=scan_path)
    pdf.save(root / "mixed-searchable-scanned.pdf")
    pdf.close()
    scan_path.unlink()

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Product", "Units"])
    summary.append(["Oranges", 42])
    archive = workbook.create_sheet("Archive")
    archive.append(["Region", "Units"])
    archive.append(["North", 7])
    workbook.save(root / "multi-sheet.xlsx")

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "Slide One Heading"
    slide.placeholders[1].text = "First slide detail"
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Slide Two Heading"
    slide.placeholders[1].text = "Second slide detail"
    presentation.save(root / "multi-slide.pptx")

    docx = root / "table.docx"
    _write_docx(docx)
    with ZipFile(root / "corrupt.docx", "w"):
        pass
    shutil.copyfile(docx, root / "macro.docm")
    (root / "unsupported.bin").write_bytes(b"unsupported fixture")

    protected = fitz.open()
    protected.new_page().insert_text((72, 72), "PRIVATE PDF")
    protected.save(
        root / "protected.pdf",
        encryption=fitz.PDF_ENCRYPT_AES_256,
        owner_pw="owner",
        user_pw="secret",
    )
    protected.close()


if __name__ == "__main__":
    generate(Path(sys.argv[1]))
