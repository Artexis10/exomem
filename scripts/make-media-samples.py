"""Write the media sample set that the Cloud image build gate extracts.

The samples live in `tests/fixtures/media-samples/`. `expected.json` maps each file
to a phrase that its extracted text must contain. The Cloud image build
(`Dockerfile`, stage `cloud-media`) and `scripts/check-media-samples.py` read it.

The script needs the document libraries, Pillow, python-docx and a CJK font. Run it
in a throwaway container, not on the host:

    docker run --rm -v "$PWD:/w" -w /w python:3.12-slim sh -c \
      "apt-get update && apt-get install -y fonts-noto-cjk fonts-dejavu-core && \
       pip install pillow pymupdf python-docx openpyxl python-pptx odfdo && \
       python scripts/make-media-samples.py tests/fixtures/media-samples"
"""

from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path

LATIN_FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
CJK_FONT = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"

# Each sample carries one phrase that search must find after extraction.
PHRASES = {
    "sample.pdf": "copper kettle quarterly almanac",
    "sample.docx": "velvet compass expedition ledger",
    "sample.xlsx": "lantern oil inventory",
    "sample.pptx": "glacier orchard roadmap",
    "sample.html": "midnight bakery opening hours",
    "sample.epub": "the cartographer of salt marshes",
    "sample.odt": "heron census at the northern marsh",
    "sample.ods": "barley harvest tally",
    "sample.odp": "tidal observatory briefing",
    "sample.rtf": "amber lighthouse keeper journal",
    "sample.txt": "pine resin storage notes",
    "sample.eml": "orchard ladder repair schedule",
    "sample.ics": "beekeeping guild assembly",
    "ocr-eng.png": "harbour lantern was repaired",
    "ocr-est.png": "Põhjatäht särab öösel",
    "ocr-jpn.png": "古い地図を見つけました",
    "ocr-jpn-vert.png": "静かな朝を過ごしました",
}

OCR_LINES = {
    "ocr-eng.png": [
        "The quiet harbour lantern was repaired",
        "before the winter storms arrived, and the",
        "fishermen thanked the old keeper warmly.",
    ],
    "ocr-est.png": [
        "Põhjatäht särab öösel vaikse järve kohal,",
        "kui külmad tuuled jõuavad mägedest",
        "üle metsade ja põldude rannani.",
    ],
    "ocr-jpn.png": [
        "昨日の午後、東京駅の近くの古本屋で",
        "古い地図を見つけました。その地図には",
        "明治時代の町並みが細かく描かれています。",
    ],
    "ocr-jpn-vert.png": [
        "京都の庭園で静かな朝を過ごしました。",
        "池の水面には紅葉が映っていました。",
        "鳥の声だけが遠くから聞こえてきます。",
    ],
}


def _png(path: Path, lines: list[str], *, font: str, vertical: bool = False) -> None:
    from PIL import Image, ImageDraw, ImageFont

    size = 44
    face = ImageFont.truetype(font, size)
    step = int(size * 1.6)
    if vertical:
        # Columns run top to bottom and are read right to left.
        height = max(len(line) for line in lines) * step + 2 * step
        width = (len(lines) + 1) * step
        image = Image.new("L", (width, height), 255)
        draw = ImageDraw.Draw(image)
        for column, line in enumerate(lines):
            x = width - (column + 1) * step
            for row, char in enumerate(line):
                draw.text((x, step + row * step), char, font=face, fill=0)
    else:
        width = max(int(draw_len) for draw_len in (face.getlength(line) for line in lines)) + 2 * step
        height = (len(lines) + 1) * step
        image = Image.new("L", (width, height), 255)
        draw = ImageDraw.Draw(image)
        for row, line in enumerate(lines):
            draw.text((step, step // 2 + row * step), line, font=face, fill=0)
    image.save(path, optimize=True)


def _epub(path: Path, phrase: str) -> None:
    chapter = (
        '<?xml version="1.0" encoding="utf-8"?>'
        '<html xmlns="http://www.w3.org/1999/xhtml"><head><title>One</title></head>'
        f"<body><h1>Chapter one</h1><p>This is the tale of {phrase}.</p></body></html>"
    )
    opf = (
        '<?xml version="1.0" encoding="utf-8"?>'
        '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id">'
        '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
        '<dc:identifier id="id">sample-epub</dc:identifier><dc:title>Sample</dc:title>'
        "<dc:language>en</dc:language></metadata>"
        '<manifest><item id="c1" href="c1.xhtml" media-type="application/xhtml+xml"/></manifest>'
        '<spine><itemref idref="c1"/></spine></package>'
    )
    container = (
        '<?xml version="1.0"?>'
        '<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
        '<rootfiles><rootfile full-path="OEBPS/content.opf" '
        'media-type="application/oebps-package+xml"/></rootfiles></container>'
    )
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        archive.writestr("META-INF/container.xml", container)
        archive.writestr("OEBPS/content.opf", opf)
        archive.writestr("OEBPS/c1.xhtml", chapter)


def main(out: Path) -> None:
    out.mkdir(parents=True, exist_ok=True)

    import docx
    import fitz
    import openpyxl
    import pptx
    from odfdo import Cell, Document, DrawPage, Frame, Header, Paragraph, Row, Table

    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), f"Notes for the {PHRASES['sample.pdf']}.")
    pdf.save(out / "sample.pdf")

    word = docx.Document()
    word.add_heading("Field notes", level=1)
    word.add_paragraph(f"Entry for the {PHRASES['sample.docx']}.")
    word.save(out / "sample.docx")

    book = openpyxl.Workbook()
    book.active.append(["item", "count"])
    book.active.append([PHRASES["sample.xlsx"], 42])
    book.save(out / "sample.xlsx")

    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Plan"
    slide.placeholders[1].text = PHRASES["sample.pptx"]
    deck.save(out / "sample.pptx")

    (out / "sample.html").write_text(
        f"<html><body><h1>Notice</h1><p>The {PHRASES['sample.html']} changed.</p></body></html>",
        encoding="utf-8",
    )
    _epub(out / "sample.epub", PHRASES["sample.epub"])

    text_doc = Document("text")
    text_doc.body.append(Header(1, "Survey"))
    text_doc.body.append(Paragraph(f"Report on the {PHRASES['sample.odt']}."))
    text_doc.save(out / "sample.odt")

    sheet_doc = Document("spreadsheet")
    table = Table("Tally")
    row = Row()
    row.set_cell(0, Cell(PHRASES["sample.ods"]))
    row.set_cell(1, Cell(7))
    table.append(row)
    sheet_doc.body.append(table)
    sheet_doc.save(out / "sample.ods")

    slides_doc = Document("presentation")
    page = DrawPage("p1", name="Slide")
    page.append(Frame.text_frame(PHRASES["sample.odp"], size=("12cm", "2cm"), position=("1cm", "1cm")))
    slides_doc.body.append(page)
    slides_doc.save(out / "sample.odp")

    (out / "sample.rtf").write_text(
        r"{\rtf1\ansi\deff0{\fonttbl{\f0 Times New Roman;}}\f0\fs24 The \b "
        + PHRASES["sample.rtf"]
        + r"\b0  begins here.\par}",
        encoding="ascii",
    )
    (out / "sample.txt").write_text(f"Reminder: {PHRASES['sample.txt']}.\n", encoding="utf-8")
    (out / "sample.eml").write_text(
        "From: keeper@example.org\nTo: crew@example.org\nSubject: Plan\n"
        "Content-Type: text/plain; charset=utf-8\n\n"
        f"Please read the {PHRASES['sample.eml']}.\n",
        encoding="utf-8",
    )
    (out / "sample.ics").write_text(
        "BEGIN:VCALENDAR\nVERSION:2.0\nBEGIN:VEVENT\nUID:sample-1\n"
        "DTSTART:20261012T090000Z\n"
        f"SUMMARY:{PHRASES['sample.ics']}\nEND:VEVENT\nEND:VCALENDAR\n",
        encoding="utf-8",
    )

    _png(out / "ocr-eng.png", OCR_LINES["ocr-eng.png"], font=LATIN_FONT)
    _png(out / "ocr-est.png", OCR_LINES["ocr-est.png"], font=LATIN_FONT)
    _png(out / "ocr-jpn.png", OCR_LINES["ocr-jpn.png"], font=CJK_FONT)
    _png(out / "ocr-jpn-vert.png", OCR_LINES["ocr-jpn-vert.png"], font=CJK_FONT, vertical=True)

    (out / "expected.json").write_text(
        json.dumps(PHRASES, ensure_ascii=False, indent=1, sort_keys=True) + "\n", encoding="utf-8"
    )


if __name__ == "__main__":
    main(Path(sys.argv[1]))
